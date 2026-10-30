#!/usr/bin/env python3
"""Advisory layout audit for an editable KMUTT-thesis DOCX.

Source: the user-supplied KMUTT 2556 postgraduate thesis manual.
No text-position or font-rendering inspection is possible from DOCX XML alone.
"""
import argparse
import json
import re
from collections import Counter
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm

EXPECTED = {'left_margin': 40, 'right_margin': 20, 'top_margin': 30, 'bottom_margin': 20}


def mm(value):
    return None if value is None else value / 36000  # python-docx EMU per mm


def run_all_paragraphs(doc):
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para


def check(filename, language):
    doc = Document(filename)
    errors, warnings, facts = [], [], []
    for idx, section in enumerate(doc.sections, 1):
        w, h = mm(section.page_width), mm(section.page_height)
        if not (abs(w-210)<0.5 and abs(h-297)<0.5):
            errors.append(f'Section {idx}: paper {w:.2f} x {h:.2f} mm, expected A4 portrait 210 x 297 mm')
        for field, exp in EXPECTED.items():
            val = mm(getattr(section, field))
            if val is None or abs(val-exp)>0.8:
                errors.append(f'Section {idx}: {field} = {val} mm, expected {exp} mm')
        facts.append(f'Section {idx}: A4 {w:.1f}x{h:.1f} mm, margins L/R/T/B '+
                     '/'.join(f'{mm(getattr(section,k)):.1f}' for k in EXPECTED)+' mm')

    style_names = {s.name for s in doc.styles}
    expected_style = 'KMUTT TH Body 16' if language=='th' else 'KMUTT EN Body 12'
    if language in ('th','en') and expected_style not in style_names:
        warnings.append(f'No optional starter style {expected_style!r}; inspect actual paragraph styles and runs manually')

    fonts, explicit_run_sizes, headings, count_text = Counter(), Counter(), [], 0
    for para in run_all_paragraphs(doc):
        txt = para.text.strip()
        if not txt:
            continue
        count_text += len(txt)
        if re.search(r'\b(รูปที่|ตารางที่)\s*\d', txt):
            headings.append(txt[:80])
        for r in para.runs:
            if r.font.name:
                fonts[r.font.name] += 1
            if r.font.size:
                explicit_run_sizes[round(r.font.size.pt,1)] += 1
    for shape_idx, shape in enumerate(doc.inline_shapes, 1):
        w, h = mm(shape.width), mm(shape.height)
        if (w > 130.8 or h > 180.8) and (h > 130.8 or w > 180.8):
            warnings.append(f'Inline image {shape_idx}: {w:.1f}x{h:.1f} mm exceeds common 130x180 mm bound')
    if count_text == 0:
        warnings.append('No document content: style-only starter is not submission-ready')
    if not doc.inline_shapes and (re.search(r'รูปที่\s*\d', '\n'.join(p.text for p in doc.paragraphs))):
        warnings.append('Figure caption pattern found, but no inline picture; could be floating image or missing figure')
    if len(doc.sections) == 1 and count_text > 10000:
        warnings.append('Single Word section with long thesis: preliminary/main/chapter-first numbering needs visual confirmation')
    if not any(s.different_first_page_header_footer for s in doc.sections) and count_text > 10000:
        warnings.append('No section has different-first-page header/footer; check if chapter opening page numbers are suppressed')
    warnings.append('Cannot prove paper weight, actual page numbers, PDF glyph correctness, exact line wrap, updated TOC, or bibliography consistency from DOCX properties')
    facts += [f'Non-empty content characters: {count_text}',
              f'Inline figures: {len(doc.inline_shapes)}',
              f'Explicit run font frequencies: {dict(fonts.most_common(8))}',
              f'Explicit run sizes: {dict(explicit_run_sizes.most_common(8))}']
    return {'path': str(Path(filename)), 'source_manual': 'KMUTT postgraduate thesis guide BE 2556',
            'errors': errors, 'warnings': warnings, 'facts': facts,
            'status': 'FAIL' if errors else 'REVIEW'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('docx')
    parser.add_argument('--language', choices=['th','en','mixed'], default='th')
    parser.add_argument('--json-out', help='Optional path for audit JSON')
    args = parser.parse_args()
    report=check(args.docx,args.language)
    print(f"KMUTT 2556 DOCX audit: {report['status']}")
    for k in ('errors','warnings','facts'):
        for x in report[k]: print(k.upper()+': '+x)
    if args.json_out:
        Path(args.json_out).write_text(json.dumps(report, ensure_ascii=False, indent=2),encoding='utf8')
    raise SystemExit(1 if report['errors'] else 0)

if __name__=='__main__': main()
